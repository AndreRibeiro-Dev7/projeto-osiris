"""Create branded Excel workbooks for financial reports."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.report import FinancialReportResponse, RevenueBreakdown

INK, GOLD, PAPER, WHITE, SOFT = "101418", "C98B3C", "F4F1EB", "FFFFFF", "697178"
CURRENCY = 'R$ #,##0.00;[Red]-R$ #,##0.00;–'


def build_financial_workbook(business_name: str, report: FinancialReportResponse) -> BytesIO:
    """Return a polished, single-sheet financial report as XLSX bytes."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Financeiro"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A15"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.tabColor = GOLD
    sheet.oddHeader.center.text = "&BProjeto Osiris — Relatório financeiro"
    sheet.oddFooter.center.text = "Página &P de &N"
    for index, width in enumerate([3, 30, 18, 18, 18, 18], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.merge_cells("B2:F2")
    sheet["B2"] = "RELATÓRIO FINANCEIRO"
    sheet["B2"].font = Font(name="Aptos Display", size=16, bold=True, color=INK)
    sheet.row_dimensions[2].height = 28
    sheet.merge_cells("B3:F3")
    sheet["B3"] = business_name
    sheet["B3"].font = Font(name="Aptos", size=11, color=SOFT)
    sheet.merge_cells("B4:F4")
    sheet["B4"] = f"Período: {report.date_from:%d/%m/%Y} a {report.date_to:%d/%m/%Y}"
    sheet["B4"].font = Font(name="Aptos", size=10, italic=True, color=SOFT)
    for column in range(2, 7):
        sheet.cell(5, column).border = Border(bottom=Side(style="thin", color=GOLD))

    cards = [
        (7, 2, "Faturamento bruto", report.total_revenue_cents / 100),
        (7, 4, "Comissões", report.total_commission_cents / 100),
        (7, 6, "Receita após comissões", report.net_revenue_cents / 100),
        (10, 2, "Despesas", report.total_expenses_cents / 100),
        (10, 4, "Resultado final", report.final_result_cents / 100),
    ]
    for row, column, label, value in cards:
        sheet.cell(row, column, label).font = Font(name="Aptos", size=9, bold=True, color=SOFT)
        cell = sheet.cell(row + 1, column, value)
        cell.number_format = CURRENCY
        cell.font = Font(name="Aptos Display", size=14, bold=True, color=INK)

    status_values = [
        ("Concluídos", report.completed), ("Agendados", report.scheduled),
        ("Confirmados", report.confirmed), ("Cancelados", report.cancelled),
        ("Não compareceram", report.no_show),
    ]
    for column, (label, value) in enumerate(status_values, start=2):
        sheet.cell(13, column, f"{label}: {value}").font = Font(name="Aptos", size=9, color=INK)

    next_row = _write_breakdown(sheet, 15, "Faturamento por profissional", report.by_barber, True)
    next_row = _write_breakdown(sheet, next_row + 2, "Faturamento por serviço", report.by_service, False)
    next_row = _write_breakdown(sheet, next_row + 2, "Faturamento por forma de pagamento", report.by_payment_method, False)
    _write_expenses(sheet, next_row + 2, report.expenses)
    sheet.print_area = f"B2:F{sheet.max_row}"
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _write_breakdown(sheet, start_row: int, title: str, rows: list[RevenueBreakdown], commissions: bool) -> int:
    sheet.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=6)
    heading = sheet.cell(start_row, 2, title)
    heading.fill = PatternFill("solid", fgColor=INK)
    heading.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    heading.alignment = Alignment(vertical="center")
    sheet.row_dimensions[start_row].height = 24
    headers = ["Nome", "Atendimentos", "Faturamento", "Comissão (%)", "Comissão"]
    if not commissions:
        headers = headers[:3]
    header_row = start_row + 1
    for column, value in enumerate(headers, start=2):
        cell = sheet.cell(header_row, column, value)
        cell.fill = PatternFill("solid", fgColor=GOLD)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    data_rows = rows or [RevenueBreakdown(label="Sem recebimentos", total_cents=0, appointments=0)]
    for row_number, item in enumerate(data_rows, start=header_row + 1):
        values = [item.label, item.appointments, item.total_cents / 100]
        if commissions:
            values.extend([(item.commission_percentage or 0) / 100, (item.commission_cents or 0) / 100])
        for column, value in enumerate(values, start=2):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_number % 2 else PAPER)
            cell.alignment = Alignment(horizontal="left" if column == 2 else "right", vertical="center")
        sheet.cell(row_number, 3).number_format = "#,##0"
        sheet.cell(row_number, 4).number_format = CURRENCY
        if commissions:
            sheet.cell(row_number, 5).number_format = "0%"
            sheet.cell(row_number, 6).number_format = CURRENCY
    return header_row + len(data_rows)


def _write_expenses(sheet, start_row: int, expenses: list) -> int:
    categories = {"rent": "Aluguel", "supplies": "Materiais", "utilities": "Contas", "marketing": "Marketing", "taxes": "Impostos", "other": "Outros"}
    sheet.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=6)
    heading = sheet.cell(start_row, 2, "Despesas do período")
    heading.fill = PatternFill("solid", fgColor=INK)
    heading.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    headers = ["Data", "Categoria", "Descrição", "Valor"]
    for column, value in enumerate(headers, start=2):
        cell = sheet.cell(start_row + 1, column, value)
        cell.fill = PatternFill("solid", fgColor=GOLD)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")
    data = expenses or [None]
    for row_number, expense in enumerate(data, start=start_row + 2):
        values = ["", "", "Sem despesas registradas", 0] if expense is None else [expense.occurred_on, categories.get(expense.category, expense.category), expense.description, expense.amount_cents / 100]
        for column, value in enumerate(values, start=2):
            cell = sheet.cell(row_number, column, value)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_number % 2 else PAPER)
            cell.font = Font(name="Aptos", size=10, color=INK)
        sheet.cell(row_number, 2).number_format = "dd/mm/yyyy"
        sheet.cell(row_number, 5).number_format = CURRENCY
    return start_row + 1 + len(data)
