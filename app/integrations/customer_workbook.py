"""Create branded Excel workbooks for the customer portfolio."""

from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.customer import CustomerPortfolioEntry

INK, GOLD, PAPER, WHITE, SOFT = "101418", "C98B3C", "F4F1EB", "FFFFFF", "697178"
CURRENCY = 'R$ #,##0.00;[Red]-R$ #,##0.00;–'


def build_customer_workbook(
    business_name: str,
    timezone: str,
    customers: list[CustomerPortfolioEntry],
) -> BytesIO:
    """Return a polished, single-sheet customer portfolio as XLSX bytes."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clientes"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B14"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.tabColor = GOLD
    sheet.oddHeader.center.text = "&BProjeto Osiris — Carteira de clientes"
    sheet.oddFooter.center.text = "Página &P de &N"

    widths = [3, 30, 20, 17, 16, 14, 12, 18, 19]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.merge_cells("B2:I2")
    sheet["B2"] = "CARTEIRA DE CLIENTES"
    sheet["B2"].font = Font(name="Aptos Display", size=16, bold=True, color=INK)
    sheet.row_dimensions[2].height = 28
    sheet.merge_cells("B3:I3")
    sheet["B3"] = business_name
    sheet["B3"].font = Font(name="Aptos", size=11, color=SOFT)
    sheet.merge_cells("B4:I4")
    generated_at = datetime.now(ZoneInfo(timezone))
    local_timezone = ZoneInfo(timezone)
    sheet["B4"] = f"Atualizado em {generated_at:%d/%m/%Y às %H:%M}"
    sheet["B4"].font = Font(name="Aptos", size=10, italic=True, color=SOFT)
    for column in range(2, 10):
        sheet.cell(5, column).border = Border(bottom=Side(style="thin", color=GOLD))

    total_completed = sum(item.completed for item in customers)
    total_no_show = sum(item.no_show for item in customers)
    total_spent = sum(item.total_spent_cents for item in customers) / 100
    cards = [
        (2, "Clientes", len(customers), "#,##0"),
        (4, "Atendimentos concluídos", total_completed, "#,##0"),
        (6, "Faltas", total_no_show, "#,##0"),
        (8, "Total gasto", total_spent, CURRENCY),
    ]
    for column, label, value, number_format in cards:
        sheet.cell(7, column, label).font = Font(
            name="Aptos", size=9, bold=True, color=SOFT
        )
        cell = sheet.cell(8, column, value)
        cell.number_format = number_format
        cell.font = Font(name="Aptos Display", size=14, bold=True, color=INK)

    sheet.merge_cells("B11:I11")
    heading = sheet["B11"]
    heading.value = "Clientes cadastrados"
    heading.fill = PatternFill("solid", fgColor=INK)
    heading.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    heading.alignment = Alignment(vertical="center")
    sheet.row_dimensions[11].height = 24

    headers = [
        "Nome",
        "Telefone",
        "Cliente desde",
        "Atendimentos",
        "Concluídos",
        "Faltas",
        "Total gasto",
        "Última visita",
    ]
    for column, value in enumerate(headers, start=2):
        cell = sheet.cell(12, column, value)
        cell.fill = PatternFill("solid", fgColor=GOLD)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    data = customers or [None]
    for row_number, customer in enumerate(data, start=13):
        values = (
            ["Nenhum cliente cadastrado", "", None, 0, 0, 0, 0, None]
            if customer is None
            else [
                customer.full_name,
                customer.phone,
                customer.created_at.astimezone(local_timezone).replace(tzinfo=None),
                customer.appointments,
                customer.completed,
                customer.no_show,
                customer.total_spent_cents / 100,
                customer.last_visit_at.astimezone(local_timezone).replace(tzinfo=None)
                if customer.last_visit_at
                else None,
            ]
        )
        for column, value in enumerate(values, start=2):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.fill = PatternFill("solid", fgColor=WHITE if row_number % 2 else PAPER)
            cell.alignment = Alignment(
                horizontal="left" if column in (2, 3) else "right",
                vertical="center",
            )
        sheet.cell(row_number, 4).number_format = "dd/mm/yyyy"
        for column in range(5, 8):
            sheet.cell(row_number, column).number_format = "#,##0"
        sheet.cell(row_number, 8).number_format = CURRENCY
        sheet.cell(row_number, 9).number_format = "dd/mm/yyyy hh:mm"

    last_row = 12 + len(data)
    sheet.auto_filter.ref = f"B12:I{last_row}"
    sheet.print_title_rows = "1:12"
    sheet.print_area = f"B2:I{last_row}"

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
