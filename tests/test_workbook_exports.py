from datetime import UTC, date, datetime
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.integrations.customer_workbook import build_customer_workbook
from app.integrations.financial_workbook import build_financial_workbook
from app.schemas.customer import CustomerPortfolioEntry
from app.schemas.expense import ExpenseResponse
from app.schemas.report import FinancialReportResponse, RevenueBreakdown


@pytest.mark.parametrize("populated", [False, True])
def test_customer_export_is_readable_with_or_without_rows(populated: bool) -> None:
    customer = CustomerPortfolioEntry(
        id=uuid4(),
        full_name="Test customer",
        phone="00000000000",
        created_at=datetime(2030, 1, 15, 14, tzinfo=UTC),
        appointments=2,
        completed=1,
        no_show=0,
        total_spent_cents=1234,
        loyalty_enabled=True,
        loyalty_progress=1,
        loyalty_target=10,
        loyalty_rewards=0,
        loyalty_rewards_redeemed=0,
        loyalty_rewards_available=0,
        loyalty_reward="Example reward",
    )

    workbook = load_workbook(
        build_customer_workbook(
            "Test business", "America/Sao_Paulo", [customer] if populated else []
        )
    )
    sheet = workbook["Clientes"]
    assert sheet["B3"].value == "Test business"
    assert sheet["B13"].value == ("Test customer" if populated else "Nenhum cliente cadastrado")
    assert sheet["H13"].value == (12.34 if populated else 0)
    if populated:
        assert sheet["D13"].value == datetime(2030, 1, 15, 11)
    workbook.close()


@pytest.mark.parametrize("populated", [False, True])
def test_financial_export_preserves_amounts_and_empty_placeholders(populated: bool) -> None:
    expense = ExpenseResponse(
        id=uuid4(),
        business_id=uuid4(),
        category="rent",
        description="Test rent",
        amount_cents=500,
        occurred_on=date(2030, 1, 15),
        is_fixed=True,
        is_paid=False,
        paid_at=None,
        created_at=datetime(2030, 1, 15, tzinfo=UTC),
    )
    row = RevenueBreakdown(label="Test cut", total_cents=1234, appointments=1)
    report = FinancialReportResponse(
        date_from=date(2030, 1, 1),
        date_to=date(2030, 1, 31),
        total_revenue_cents=1234 if populated else 0,
        total_commission_cents=0,
        net_revenue_cents=1234 if populated else 0,
        total_expenses_cents=500 if populated else 0,
        final_result_cents=734 if populated else 0,
        average_ticket_cents=1234 if populated else 0,
        attendance_rate_percent=100 if populated else 0,
        completed=int(populated),
        scheduled=0,
        confirmed=0,
        cancelled=0,
        no_show=0,
        by_barber=[row] if populated else [],
        by_service=[row] if populated else [],
        by_payment_method=[row] if populated else [],
        expenses=[expense] if populated else [],
        daily=[],
    )

    workbook = load_workbook(build_financial_workbook("Test business", report))
    sheet = workbook["Financeiro"]
    assert sheet["B8"].value == (12.34 if populated else 0)
    assert sheet["B17"].value == ("Test cut" if populated else "Sem recebimentos")
    assert sheet["D17"].value == (12.34 if populated else 0)
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert ("Test rent" if populated else "Sem despesas registradas") in values
    workbook.close()
